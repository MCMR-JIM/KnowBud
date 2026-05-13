from __future__ import annotations

import json
import logging
import os
import re
import subprocess
import sys
import tempfile
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Callable

from src.core.models import (
    CandidateEdge,
    CandidateGraph,
    CandidateNode,
    CandidateRelay,
    DocumentMeta,
    ExtractionDiagnostics,
    ResourceRecord,
    ResourceSegment,
    SourceRef,
    SourceSegment,
    TopicNode,
)
from src.agent.models import EdgeType
from src.services.llm_gateway import create_chat_completion, extract_message_text
from src.services.resource_graph_curation import (
    _subject_from_tags,
    create_resource_level_proposals,
    detect_resource_subject,
)

if TYPE_CHECKING:
    from src.services.session_backend import SessionBackend


TARGET_CHARS = 1000
MAX_CHARS = 1800
CLASSIFIED_THRESHOLD = 0.45
CLASSIFY_BATCH_SIZE = 1

logger = logging.getLogger(__name__)


def _llm_call(backend, **kwargs):
    backend.llm_skill.ensure_configured()
    client = getattr(backend.llm_skill, "client", None)
    if client is None:
        raise RuntimeError("LLM client is not configured")
    return create_chat_completion(
        client=client,
        base_url=getattr(backend.llm_skill, "base_url", ""),
        model_name=kwargs.get("model") or getattr(backend.llm_skill, "model_name", ""),
        messages=kwargs.get("messages") or [],
        temperature=float(kwargs.get("temperature", 0.3)),
        timeout=float(kwargs.get("timeout", 60.0)),
        extra_body=kwargs.get("extra_body"),
    )


def _console_log(title: str, payload: object) -> None:
    try:
        if isinstance(payload, (dict, list)):
            rendered = json.dumps(payload, ensure_ascii=False, indent=2)
        else:
            rendered = str(payload)
    except Exception:
        rendered = str(payload)
    print(f"\n[{title}]\n{rendered}\n", flush=True)


def _strip_images(text: str) -> str:
    import re as _re
    return _re.sub(r"!\[.*?\]\(.*?\)", "", text).strip()


def _safe_progress(
    progress_callback: Callable[[str, dict[str, object]], None] | None,
    event: str,
    payload: dict[str, object],
) -> None:
    if progress_callback is None:
        return
    try:
        progress_callback(event, payload)
    except Exception:
        logger.exception("progress callback failed", extra={"event": event})


def _estimate_tokens(text: str) -> int:
    cjk = sum(1 for c in text if '\u4e00' <= c <= '\u9fff' or '\u3400' <= c <= '\u4dbf')
    return int(cjk * 1.5 + (len(text) - cjk) * 0.3)


@dataclass
class TextUnit:
    text: str
    locator: dict[str, object]
    mineru_block: dict | None = None


def ingest_document_resource(
    backend: "SessionBackend",
    record: ResourceRecord,
    topics: list[TopicNode],
    default_topic_id: str | None = None,
    progress_callback: Callable[[str, dict[str, object]], None] | None = None,
    enable_graph_search: bool = False,
    enable_new_pipeline: bool = False,
    subject: str | None = None,
    language_id: str | None = None,
    skip_local_parsing: bool = False,
) -> list[ResourceSegment]:
    import os as _os

    units: list[TextUnit] | None = None
    parser = _select_parser(record)
    allow_mineru = (
        _os.getenv("MINERU_ENABLED", "true").strip().lower() != "false"
        and not skip_local_parsing  # 远程 API 直传模式：跳过本地 ML 模型解析
    )

    if allow_mineru:
        _mineru_result: dict | None = None
        try:
            _lang = _mineru_lang_code(language_id)
            _mineru_result = _parse_with_mineru(str(record.stored_path), lang=_lang)
            if _mineru_result["error"] is None:
                units = _convert_mineru_to_textunits(_mineru_result)
            else:
                logger.warning("MinerU parse failed, falling back: %s", _mineru_result["error"])
        except Exception as _mineru_exc:
            logger.warning("MinerU invocation failed, falling back: %s", str(_mineru_exc))
        finally:
            if _mineru_result and _mineru_result.get("images_dir"):
                _cleanup_mineru_temp(_mineru_result["images_dir"])

    if units is None:
        if parser is None:
            segments = [
                _status_segment(
                    resource_id=record.resource_id,
                    sequence_index=0,
                    status="unsupported",
                    reason=f"暂不支持解析 {record.media_type} 文件",
                    locator={"kind": record.media_type or "binary"},
                )
            ]
            backend.replace_resource_segments(record.resource_id, segments)
            return segments

        try:
            units = parser(Path(record.stored_path))
        except Exception as exc:
            segments = [
                _status_segment(
                    resource_id=record.resource_id,
                    sequence_index=0,
                    status="parse_failed",
                    reason=f"文档解析失败: {str(exc)[:60]}",
                    locator={"kind": record.media_type or "unknown"},
                )
            ]
            backend.replace_resource_segments(record.resource_id, segments)
            return segments

    chunks = _units_to_chunks(units)
    _safe_progress(
        progress_callback,
        "document_parsed",
        {
            "resource_id": record.resource_id,
            "chunk_count": len(chunks),
            "media_type": record.media_type,
        },
    )
    if not chunks:
        segments = [
            _status_segment(
                resource_id=record.resource_id,
                sequence_index=0,
                status="parse_failed",
                reason="文档未提取到可用文本",
                locator={"kind": record.media_type or "unknown"},
            )
        ]
        backend.replace_resource_segments(record.resource_id, segments)
        return segments

    if enable_new_pipeline and subject:
        segments = _run_structured_ingestion(
            backend=backend, record=record, chunks=chunks,
            topics=topics, subject=subject, language_id=language_id,
            progress_callback=progress_callback,
        )
        backend.replace_resource_segments(record.resource_id, segments)
        return segments

    _ctx_subject: str | None = None
    _ctx_language_id: str | None = None
    if enable_graph_search:
        try:
            detected = detect_resource_subject(
                record=record,
                segments=[],
                topics=topics,
                default_topic_id=default_topic_id or record.topic_id,
                backend=backend,
            )
            if detected != "general":
                _ctx_subject = detected
        except Exception:
            _ctx_subject = None

    segments: list[ResourceSegment] = []
    batch_size = 1 if len(chunks) <= 1 else CLASSIFY_BATCH_SIZE
    for batch_start in range(0, len(chunks), batch_size):
        batch = chunks[batch_start:batch_start + batch_size]
        if batch_size == 1:
            segments.append(
                _classify_chunk(
                    backend=backend,
                    resource_id=record.resource_id,
                    sequence_index=batch_start,
                    chunk=batch[0],
                    topics=topics,
                    default_topic_id=default_topic_id or record.topic_id,
                    subject=_ctx_subject,
                    language_id=_ctx_language_id,
                )
            )
        else:
            batch_segments = _classify_chunks_batch(
                backend=backend,
                resource_id=record.resource_id,
                start_index=batch_start,
                chunks=batch,
                topics=topics,
                default_topic_id=default_topic_id or record.topic_id,
                subject=_ctx_subject,
                language_id=_ctx_language_id,
            )
            segments.extend(batch_segments)

    create_resource_level_proposals(
        backend=backend,
        record=record,
        segments=segments,
        topics=topics,
        default_topic_id=default_topic_id or record.topic_id,
    )

    backend.replace_resource_segments(record.resource_id, segments)
    return segments


def _select_parser(record: ResourceRecord) -> Callable[[Path], list[TextUnit]] | None:
    media_type = (record.media_type or "").strip().lower()
    ext = Path(record.original_filename).suffix.lower()
    if media_type == "txt" or ext in {".txt", ".md"}:
        return _parse_txt
    if media_type == "pdf" or ext == ".pdf":
        return _parse_pdf
    if media_type == "docx" or ext == ".docx":
        return _parse_docx
    if media_type == "pptx" or ext == ".pptx":
        return _parse_pptx
    if media_type == "doc" or ext == ".doc":
        return None
    return None


def _parse_txt(path: Path) -> list[TextUnit]:
    text = path.read_text(encoding="utf-8-sig", errors="ignore")
    lines = text.splitlines()
    units: list[TextUnit] = []
    buffer: list[str] = []
    line_start: int | None = None
    for index, raw_line in enumerate(lines, start=1):
        stripped = raw_line.strip()
        if stripped:
            if line_start is None:
                line_start = index
            buffer.append(stripped)
            continue
        if buffer and line_start is not None:
            units.append(
                TextUnit(
                    text="\n".join(buffer),
                    locator={"kind": "txt", "line_start": line_start, "line_end": index - 1},
                )
            )
            buffer = []
            line_start = None
    if buffer and line_start is not None:
        units.append(
            TextUnit(
                text="\n".join(buffer),
                locator={"kind": "txt", "line_start": line_start, "line_end": len(lines) or line_start},
            )
        )
    return units


def _parse_pdf(path: Path) -> list[TextUnit]:
    try:
        import fitz  
        from rapidocr_onnxruntime import RapidOCR
    except ModuleNotFoundError as exc:
        raise RuntimeError("需要高级PDF和OCR解析库，请先执行: pip install pymupdf rapidocr-onnxruntime") from exc

    # 初始化轻量级、极其稳定的 ONNX 引擎 OCR
    ocr = RapidOCR()

    doc = fitz.open(str(path))
    units: list[TextUnit] = []

    for page_index, page in enumerate(doc, start=1):
        # 1. 尝试提取原生文本
        text = page.get_text().strip()

        # 2. 如果没字（说明是扫描全能王的图片），启动 OCR
        if not text:
            # 渲染成两倍高清图片字节流
            pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
            img_bytes = pix.tobytes("png")
            
            # 直接把图片流喂给 RapidOCR，简单粗暴不出错
            result, _ = ocr(img_bytes)
            
            if result:
                # 提取识别出的文字
                text = "\n".join([line[1] for line in result])

        text = text.strip()
        if text:
            for block in _split_text_blocks(text):
                units.append(
                    TextUnit(
                        text=block,
                        locator={"kind": "pdf", "page_start": page_index, "page_end": page_index},
                    )
                )
                
    doc.close()
    return units


def _parse_docx(path: Path) -> list[TextUnit]:
    try:
        from docx import Document
    except ModuleNotFoundError as exc:
        raise RuntimeError("python-docx is required for DOCX parsing") from exc

    document = Document(str(path))
    units: list[TextUnit] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        units.append(
            TextUnit(
                text=text,
                locator={"kind": "docx", "paragraph_start": paragraph_index, "paragraph_end": paragraph_index},
            )
        )
    return units


def _parse_pptx(path: Path) -> list[TextUnit]:
    try:
        from pptx import Presentation
    except ModuleNotFoundError as exc:
        raise RuntimeError("python-pptx is required for PPTX parsing") from exc

    presentation = Presentation(str(path))
    units: list[TextUnit] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if isinstance(text, str) and text.strip():
                parts.append(text.strip())
        merged = "\n".join(parts).strip()
        if not merged:
            continue
        units.append(
            TextUnit(
                text=merged,
                locator={"kind": "pptx", "slide_start": slide_index, "slide_end": slide_index},
            )
        )
    return units


def _cleanup_mineru_temp(dir_path: str) -> None:
    from src.core.safe_delete import SafeDelete
    SafeDelete(dir_path).require_in_temp_dir().require_name_starts_with("mineru_").execute(ignore_errors=True)


def _mineru_lang_code(language_id: str | None) -> str:
    if not language_id:
        return "ch"
    lid = language_id.strip().lower()
    if lid in ("en", "english", "eng"):
        return "en"
    if lid in ("zh", "chinese", "chi", "ch", "cn"):
        return "ch"
    return "ch"


def _parse_with_mineru(file_path: str, lang: str = "ch") -> dict:
    result: dict = {
        "markdown": "",
        "content_list": [],
        "images_dir": None,
        "error": None,
    }
    worker = Path(__file__).resolve().parents[2] / "scripts" / "mineru_worker.py"
    with tempfile.NamedTemporaryFile(prefix="mineru_result_", suffix=".json", delete=False) as handle:
        output_path = Path(handle.name)

    try:
        project_root = Path(__file__).resolve().parents[2]
        env = os.environ.copy()
        existing_python_path = env.get("PYTHONPATH", "")
        env["PYTHONPATH"] = str(project_root) if not existing_python_path else f"{project_root}{os.pathsep}{existing_python_path}"
        command = [
            sys.executable,
            str(worker),
            "--file",
            file_path,
            "--lang",
            lang,
            "--output",
            str(output_path),
        ]
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=1800,
            check=False,
            cwd=str(project_root),
            env=env,
        )
        if completed.returncode != 0:
            stderr = (completed.stderr or completed.stdout or "").strip()
            result["error"] = stderr[:500] or f"MinerU worker failed with exit code {completed.returncode}"
            return result
        if not output_path.exists():
            result["error"] = "MinerU worker did not produce output"
            return result
        payload = json.loads(output_path.read_text(encoding="utf-8"))
        if isinstance(payload, dict):
            result.update(payload)
        else:
            result["error"] = "MinerU worker returned invalid payload"
    except subprocess.TimeoutExpired:
        result["error"] = "MinerU parsing timed out"
    except Exception as exc:
        result["error"] = str(exc)
    finally:
        try:
            output_path.unlink(missing_ok=True)
        except Exception:
            pass
    return result


def _mineru_parse_office(file_bytes: bytes, suffix: str) -> dict:
    import tempfile
    from mineru.data.data_reader_writer import FileBasedDataWriter
    from mineru.backend.office.office_middle_json_mkcontent import union_make as office_union_make
    from mineru.utils.enum_class import MakeMode

    if suffix == ".docx":
        from mineru.backend.office.docx_analyze import office_docx_analyze
        analyze = office_docx_analyze
    else:
        from mineru.backend.office.pptx_analyze import office_pptx_analyze
        analyze = office_pptx_analyze

    images_dir = tempfile.mkdtemp(prefix="mineru_office_")
    try:
        image_writer = FileBasedDataWriter(images_dir)
        middle_json, _infer_result = analyze(file_bytes, image_writer=image_writer)
        pdf_info = middle_json["pdf_info"]
        image_dir_name = Path(images_dir).name

        markdown = office_union_make(pdf_info, MakeMode.MM_MD, image_dir_name)
        content_list = office_union_make(pdf_info, MakeMode.CONTENT_LIST_V2, image_dir_name)

        return {
            "markdown": markdown if isinstance(markdown, str) else str(markdown),
            "content_list": content_list if isinstance(content_list, list) else [],
            "images_dir": images_dir,
            "error": None,
        }
    except Exception:
        from src.core.safe_delete import SafeDelete
        SafeDelete(images_dir).require_in_temp_dir().require_name_starts_with("mineru_").execute(ignore_errors=True)
        raise


def _mineru_parse_pdf_image(file_bytes: bytes, suffix: str, stem: str, lang: str = "ch") -> dict:
    import tempfile
    from mineru.cli.common import do_parse
    from mineru.utils.enum_class import MakeMode

    if suffix in (".png", ".jpg", ".jpeg"):
        from mineru.utils.pdf_image_tools import images_bytes_to_pdf_bytes
        file_bytes = images_bytes_to_pdf_bytes(file_bytes)

    output_dir = tempfile.mkdtemp(prefix="mineru_pdf_")
    try:
        pdf_file_name = stem or "document"
        do_parse(
            output_dir=output_dir,
            pdf_file_names=[pdf_file_name],
            pdf_bytes_list=[file_bytes],
            p_lang_list=[lang],
            backend="pipeline",
            f_dump_md=True,
            f_dump_content_list=True,
            f_dump_middle_json=False,
            f_dump_model_output=False,
            f_dump_orig_pdf=False,
            f_make_md_mode=MakeMode.MM_MD,
            f_draw_layout_bbox=False,
            f_draw_span_bbox=False,
        )

        md_files = list(Path(output_dir).rglob(f"*/{pdf_file_name}.md"))
        content_list_files = list(Path(output_dir).rglob(f"*/{pdf_file_name}_content_list_v2.json"))

        markdown = ""
        content_list: list = []

        if md_files:
            markdown = md_files[0].read_text(encoding="utf-8", errors="ignore")
        if content_list_files:
            content_list = json.loads(content_list_files[0].read_text(encoding="utf-8", errors="ignore"))
            if not isinstance(content_list, list):
                content_list = []

        return {
            "markdown": markdown,
            "content_list": content_list,
            "images_dir": output_dir,
            "error": None,
        }
    except Exception:
        from src.core.safe_delete import SafeDelete
        SafeDelete(output_dir).require_in_temp_dir().require_name_starts_with("mineru_").execute(ignore_errors=True)
        raise


def _convert_mineru_to_textunits(mineru_result: dict) -> list[TextUnit]:
    markdown = mineru_result.get("markdown") or ""
    if not markdown:
        content_list = mineru_result.get("content_list") or []
        if isinstance(content_list, list) and content_list:
            return _convert_mineru_content_list_to_textunits(content_list)
        return []

    # Split markdown by headings into sections, keeping image links intact
    import re as _re
    sections = _re.split(r"\n(?=#)", markdown)
    units: list[TextUnit] = []
    heading_stack: list[str] = []

    for section in sections:
        section = section.strip()
        if not section:
            continue
        lines = section.split("\n")
        title = lines[0].strip("# ").strip() if lines[0].startswith("#") else ""
        body = "\n".join(lines[1:]).strip() if title else section

        if title:
            # Update heading stack
            level = len(lines[0]) - len(lines[0].lstrip("#"))
            while len(heading_stack) >= level:
                heading_stack.pop()
            heading_stack.append(title)

        # Create one TextUnit per section with full markdown
        text = section
        if len(text) > MAX_CHARS * 3:
            # Split very large sections on double newlines
            for para in _re.split(r"\n\n+", body):
                para = para.strip()
                if not para:
                    continue
                units.append(TextUnit(
                    text=para,
                    locator={"kind": "mineru", "heading_path": list(heading_stack)},
                    mineru_block={"section_title": title},
                ))
        else:
            units.append(TextUnit(
                text=text,
                locator={"kind": "mineru", "heading_path": list(heading_stack)},
                mineru_block={"section_title": title},
            ))
    return units


def _convert_mineru_content_list_to_textunits(content_list: list) -> list[TextUnit]:
    """Fallback: convert content_list_v2 to TextUnits (original logic)."""
    units: list[TextUnit] = []
    heading_stack: list[str] = []
    for page_idx, page_blocks in enumerate(content_list):
        if not isinstance(page_blocks, list):
            continue
        for block in page_blocks:
            if not isinstance(block, dict):
                continue
            block_type = (block.get("type") or "").lower()
            content = block.get("content") or {}
            if block_type in ("page_header", "page_footer", "page_number", "page_aside_text", "page_footnote", "image"):
                continue
            if block_type == "title":
                text = _mineru_extract_text((content.get("title_content") or content.get("content") or []))
                if not text:
                    continue
                level = content.get("level", 1)
                if isinstance(level, (int, float)):
                    level = int(level)
                else:
                    level = 1
                while len(heading_stack) >= level:
                    heading_stack.pop()
                heading_stack.append(text)
                units.append(TextUnit(text=text, locator={"kind": "mineru", "page_idx": page_idx, "block_type": "title", "heading_path": list(heading_stack[:-1])}, mineru_block=block))
            elif block_type == "paragraph":
                text = _mineru_extract_text((content.get("paragraph_content") or content.get("content") or []))
                if text:
                    units.append(TextUnit(text=text, locator={"kind": "mineru", "page_idx": page_idx, "block_type": "paragraph", "heading_path": list(heading_stack)}, mineru_block=block))
            elif block_type == "table":
                html = content.get("html") or ""
                md_table = _html_table_to_markdown(html) if html else ""
                if md_table:
                    units.append(TextUnit(text=md_table, locator={"kind": "mineru", "page_idx": page_idx, "block_type": "table", "heading_path": list(heading_stack)}, mineru_block=block))
            elif block_type == "list":
                list_items = content.get("list_items") or []
                parts = [item.get("item_content", []) for item in list_items if isinstance(item, dict)]
                text = "\n".join(f"- {_mineru_extract_text(p)}" for p in parts if _mineru_extract_text(p))
                if text:
                    units.append(TextUnit(text=text, locator={"kind": "mineru", "page_idx": page_idx, "block_type": "list", "heading_path": list(heading_stack)}, mineru_block=block))
    return units


def _mineru_extract_text(spans: list) -> str:
    parts: list[str] = []
    for span in spans:
        if not isinstance(span, dict):
            continue
        span_type = span.get("type", "")
        span_text = span.get("content", "")
        if not isinstance(span_text, str):
            continue
        if span_type in ("text", "md", ""):
            parts.append(span_text)
        elif span_type == "equation_inline":
            parts.append(f"${span_text}$")
        elif span_type == "code_inline":
            parts.append(f"`{span_text}`")
        elif span_type == "phonetic":
            parts.append(f"[{span_text}]")
    return "".join(parts).strip()


def _html_table_to_markdown(html: str) -> str:
    from html.parser import HTMLParser

    class _TableParser(HTMLParser):
        def __init__(self):
            super().__init__()
            self.rows: list[list[str]] = []
            self._current_row: list[str] = []
            self._current_cell: str = ""
            self._in_cell = False
            self._colspan = 1

        def handle_starttag(self, tag: str, attrs: list):
            tag_lower = tag.lower()
            if tag_lower in ("td", "th"):
                self._in_cell = True
                self._current_cell = ""
                for k, v in attrs:
                    if k.lower() == "colspan":
                        try:
                            self._colspan = int(v)
                        except ValueError:
                            self._colspan = 1

        def handle_endtag(self, tag: str):
            tag_lower = tag.lower()
            if tag_lower in ("td", "th"):
                self._in_cell = False
                cell = self._current_cell.strip().replace("|", "\\|").replace("\n", " ")
                for _ in range(self._colspan):
                    self._current_row.append(cell)
                self._colspan = 1
            elif tag_lower == "tr":
                if self._current_row:
                    self.rows.append(self._current_row)
                self._current_row = []

        def handle_data(self, data: str):
            if self._in_cell:
                self._current_cell += data

    parser = _TableParser()
    try:
        parser.feed(html)
    except Exception:
        return ""

    rows = parser.rows
    if not rows:
        return ""

    col_count = max((len(r) for r in rows), default=0)
    if col_count == 0:
        return ""

    lines: list[str] = []
    for ri, row in enumerate(rows):
        padded = row + [""] * (col_count - len(row))
        lines.append("| " + " | ".join(padded) + " |")
        if ri == 0:
            lines.append("| " + " | ".join(["---"] * col_count) + " |")

    return "\n".join(lines)


def _split_text_blocks(text: str) -> list[str]:
    cleaned = text.strip()
    if not cleaned:
        return []
    blocks = [block.strip() for block in re.split(r"\n\s*\n+", cleaned) if block.strip()]
    return blocks or [cleaned]


def _units_to_chunks(units: list[TextUnit]) -> list[TextUnit]:
    chunks: list[TextUnit] = []
    for unit in units:
        chunks.extend(_split_unit(unit))
    return chunks


def _split_unit(unit: TextUnit) -> list[TextUnit]:
    text = unit.text.strip()
    if not text:
        return []
    if len(text) <= MAX_CHARS:
        return [TextUnit(text=text, locator=dict(unit.locator))]

    chunks: list[TextUnit] = []
    current = ""
    for sentence in _split_sentences(text):
        if not sentence:
            continue
        if current and len(current) + len(sentence) > MAX_CHARS:
            chunks.append(TextUnit(text=current.strip(), locator=dict(unit.locator)))
            current = ""
        if len(sentence) > MAX_CHARS:
            if current:
                chunks.append(TextUnit(text=current.strip(), locator=dict(unit.locator)))
                current = ""
            for piece in _hard_split(sentence):
                chunks.append(TextUnit(text=piece, locator=dict(unit.locator)))
            continue
        if current and len(current) + len(sentence) > TARGET_CHARS:
            chunks.append(TextUnit(text=current.strip(), locator=dict(unit.locator)))
            current = sentence
        else:
            current += sentence
    if current.strip():
        chunks.append(TextUnit(text=current.strip(), locator=dict(unit.locator)))
    return chunks


def _split_sentences(text: str) -> list[str]:
    parts = re.split(r"(?<=[。！？!?；;.!?])|\n+", text)
    return [part.strip() for part in parts if part and part.strip()]


def _hard_split(text: str) -> list[str]:
    pieces: list[str] = []
    remaining = text.strip()
    while remaining:
        if len(remaining) <= MAX_CHARS:
            pieces.append(remaining)
            break
        cut = remaining[:MAX_CHARS]
        boundary = max(cut.rfind(marker) for marker in ("。", "！", "？", "；", ";", ".", "，", ",", " "))
        if boundary < int(MAX_CHARS * 0.5):
            boundary = MAX_CHARS
        else:
            boundary += 1
        pieces.append(remaining[:boundary].strip())
        remaining = remaining[boundary:].strip()
    return [piece for piece in pieces if piece]


def _classify_chunk(
    *,
    backend: "SessionBackend",
    resource_id: str,
    sequence_index: int,
    chunk: TextUnit,
    topics: list[TopicNode],
    default_topic_id: str | None,
    subject: str | None = None,
    language_id: str | None = None,
) -> ResourceSegment:
    text = chunk.text.strip()
    normalized = {
        "decision": "unclassified",
        "topic_id": None,
        "confidence": 0.0,
        "reason": "未找到合适知识点",
        "proposed_topic": None,
        "guiding_question": "",
        "teaching_hint": "",
    }
    proposed_topic_title: str | None = None
    proposed_parent_node_ids: list[str] = []

    try:
        result = backend.llm_skill.classify_or_propose_resource_chunk(
            chunk_text=_strip_images(text),
            topics=topics,
            default_parent_topic_id=default_topic_id,
        )
        if isinstance(result, dict):
            normalized = _normalize_chunk_decision(result=result, topics=topics, default_topic_id=default_topic_id)
            if normalized["decision"] == "propose":
                proposed_topic = normalized["proposed_topic"]
                if isinstance(proposed_topic, dict):
                    proposed_topic_title = str(proposed_topic.get("title", "")).strip() or None
                    proposed_parent_node_ids = proposed_topic.get("parent_node_ids") or []
    except Exception as exc:
        normalized = {
            "decision": "unclassified",
            "topic_id": None,
            "confidence": 0.0,
            "reason": f"分类失败: {str(exc)[:60]}",
            "proposed_topic": None,
            "guiding_question": "",
            "teaching_hint": "",
        }

    if normalized["decision"] == "propose" and subject is not None and backend.llm_skill.client.api_key:
        try:
            from src.services.graph_search import GraphSearchAgent

            agent = GraphSearchAgent(backend, subject, language_id)
            search_result = agent.search(chunk_text=text)
            if search_result.decision == "link" and search_result.confidence >= 0.55:
                normalized["decision"] = "link"
                normalized["topic_id"] = search_result.matched_topic_id
                normalized["confidence"] = max(normalized["confidence"], search_result.confidence)
                normalized["reason"] = f"[图谱去重] {search_result.reason[:60]}"
                normalized["proposed_topic"] = None
                proposed_topic_title = None
                proposed_parent_node_ids = []
            elif search_result.decision == "propose":
                if search_result.proposed_title:
                    proposed_topic_title = search_result.proposed_title
                if search_result.parent_node_ids:
                    proposed_parent_node_ids = search_result.parent_node_ids
                if proposed_topic_title:
                    try:
                        from src.services.resource_graph_curation import deduplicate_candidate_title

                        state = backend.load_app_state(include_history=False)
                        existing = [
                            {
                                "topic_id": t.topic_id,
                                "title": t.title,
                                "subject": _node_subject_tag(t.tags),
                                "facet": _tag_value(t.tags, "facet:"),
                                "language_id": _tag_value(t.tags, "language:"),
                                "tags": t.tags,
                            }
                            for t in state.curriculum.topics
                        ]
                        for p in state.learning.graph_proposals:
                            if p.status in {"rejected"}:
                                continue
                            existing.append(
                                {
                                    "topic_id": p.proposal_id,
                                    "title": p.title,
                                    "subject": _node_subject_tag(getattr(p, "tags", [])),
                                    "facet": _tag_value(getattr(p, "tags", []), "facet:"),
                                    "language_id": _tag_value(getattr(p, "tags", []), "language:"),
                                    "tags": getattr(p, "tags", []),
                                }
                            )
                        matched_id = deduplicate_candidate_title(
                            backend=backend,
                            candidate_title=proposed_topic_title,
                            candidate_subject=subject,
                            candidate_facet="",
                            candidate_text=text,
                            candidate_language_id=language_id,
                            existing_nodes=existing,
                        )
                        if matched_id:
                            if isinstance(matched_id, str) and matched_id.startswith("proposal_"):
                                normalized["decision"] = "propose"
                                normalized["topic_id"] = None
                                normalized["proposal_id"] = matched_id
                            else:
                                normalized["decision"] = "link"
                                normalized["topic_id"] = matched_id
                            normalized["confidence"] = max(normalized["confidence"], 0.8)
                            normalized["reason"] = f"[语义消歧] 归并到已有节点 {matched_id}"
                            normalized["proposed_topic"] = None
                            proposed_topic_title = None
                            proposed_parent_node_ids = []
                    except Exception:
                        pass
        except Exception:
            pass

    if normalized["decision"] == "link" and normalized["topic_id"]:
        status = "classified"
    elif normalized.get("proposal_id") and normalized["decision"] == "propose":
        status = "proposed"
    elif normalized["decision"] == "propose" and proposed_topic_title:
        status = "unclassified"
    else:
        status = "unclassified"
        proposed_topic_title = None

    return ResourceSegment(
        segment_id=_segment_id(resource_id, sequence_index),
        start_ms=0,
        end_ms=None,
        label="chunk",
        status=status,
        sequence_index=sequence_index,
        text=text,
        locator=dict(chunk.locator),
        topic_id=str(normalized["topic_id"]) if normalized["topic_id"] is not None else None,
        proposal_id=str(normalized.get("proposal_id")) if normalized.get("proposal_id") is not None else None,
        proposed_topic_title=proposed_topic_title,
        proposed_parent_node_ids=proposed_parent_node_ids,
        decision=str(normalized["decision"]),
        confidence=float(normalized["confidence"]),
        reason=str(normalized["reason"]),
        guiding_question=str(normalized.get("guiding_question", "")) or None,
        teaching_hint=str(normalized.get("teaching_hint", "")) or None,
    )


def _classify_chunks_batch(
    *,
    backend: "SessionBackend",
    resource_id: str,
    start_index: int,
    chunks: list[TextUnit],
    topics: list[TopicNode],
    default_topic_id: str | None,
    subject: str | None = None,
    language_id: str | None = None,
) -> list[ResourceSegment]:
    import json as _json

    topic_payload = [
        {
            "topic_id": t.topic_id,
            "title": t.title,
            "tags": t.tags,
            "prerequisite_ids": t.prerequisite_ids,
        }
        for t in topics
    ]

    chunk_texts = []
    for i, chunk in enumerate(chunks):
        chunk_texts.append(f"片段 {i}:\n```text\n{_strip_images(chunk.text.strip())[:2000]}\n```\n")
    combined = "\n".join(chunk_texts)

    system = (
        "你是儿童学习知识图谱策展助手。你的任务是判断多个教学资源片段各自应该挂到已有知识节点，"
        "还是应该提出一个新知识节点。\n"
        "返回一个 JSON 数组，每个元素对应一个片段，顺序与输入一致。\n"
        "【节点合并与复用规则】\n"
        "1. 只能复用给定的已有节点；如果没有合适节点，但片段表达了明确、可教学的知识点，请提出新节点。\n"
        "2. 如果你要提出的新节点和已有节点的概念本质相同（只是换了种说法），请直接选择 link 到已有节点！\n"
        "3. 不要为了泛泛内容创建节点。\n"
        "【新节点命名规范】\n"
        "1. 提取最精炼的专业术语或核心概念，作为新知识点的名称。\n"
        "2. 绝对不要包含「什么是」、「关于」、「浅析」等冗余前缀。\n"
        "3. 绝对不要包含「是什么」、「及其应用」、「的证明」等冗余后缀。\n"
    )
    user = (
        f"已有知识节点列表：\n{_json.dumps(topic_payload, ensure_ascii=False)}\n\n"
        f"默认父节点：\n{_json.dumps(default_topic_id, ensure_ascii=False)}\n\n"
        f"{combined}"
        "返回 JSON 数组：\n"
        '[{"decision":"link|propose|unclassified","topic_id":string|null,'
        '"confidence":number,"reason":string,"proposed_topic":{"title":string,'
        '"summary":string,"parent_node_ids":[string],"edge_type":"requires"|...}|null,'
        '"guiding_question":string,"teaching_hint":string}]'
    )

    raw_results: list[dict] = []
    try:
        response = _llm_call(backend, 
            model=backend.llm_skill.model_name,
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
            temperature=0.3,
            timeout=600.0,
            extra_body={"thinking": {"type": "disabled"}} if "deepseek" in (getattr(backend.llm_skill, "base_url", "") or "") else None,
        )
        content = (response.choices[0].message.content or "").strip()
        content = re.sub(r"^```(?:json)?\s*", "", content)
        content = re.sub(r"\s*```$", "", content)
        parsed = _json.loads(content)
        if isinstance(parsed, list):
            raw_results = [item for item in parsed if isinstance(item, dict)]
    except Exception:
        raw_results = []

    segments: list[ResourceSegment] = []
    for i, chunk in enumerate(chunks):
        seg_index = start_index + i
        if i < len(raw_results):
            result = raw_results[i]
            normalized = _normalize_chunk_decision(
                result=result, topics=topics,
                default_topic_id=default_topic_id,
            )
        else:
            normalized = {
                "decision": "unclassified",
                "topic_id": None,
                "confidence": 0.0,
                "reason": "batch classify missed this chunk",
                "proposed_topic": None,
                "guiding_question": "",
                "teaching_hint": "",
            }

        proposed_topic_title: str | None = None
        proposed_parent_node_ids: list[str] = []
        if normalized["decision"] == "propose":
            pt = normalized["proposed_topic"]
            if isinstance(pt, dict):
                proposed_topic_title = str(pt.get("title", "")).strip() or None
                proposed_parent_node_ids = pt.get("parent_node_ids") or []

        if normalized["decision"] == "link" and normalized["topic_id"]:
            status = "classified"
        elif normalized.get("proposal_id") and normalized["decision"] == "propose":
            status = "proposed"
        elif normalized["decision"] == "propose" and proposed_topic_title:
            status = "unclassified"
        else:
            status = "unclassified"
            proposed_topic_title = None

        segments.append(
            ResourceSegment(
                segment_id=_segment_id(resource_id, seg_index),
                start_ms=0,
                end_ms=None,
                label="chunk",
                status=status,
                sequence_index=seg_index,
                text=chunk.text.strip(),
                locator=dict(chunk.locator),
                topic_id=str(normalized["topic_id"]) if normalized["topic_id"] is not None else None,
                proposal_id=str(normalized.get("proposal_id")) if normalized.get("proposal_id") is not None else None,
                proposed_topic_title=proposed_topic_title,
                proposed_parent_node_ids=proposed_parent_node_ids,
                decision=str(normalized["decision"]),
                confidence=float(normalized["confidence"]),
                reason=str(normalized["reason"]),
                guiding_question=str(normalized.get("guiding_question", "")) or None,
                teaching_hint=str(normalized.get("teaching_hint", "")) or None,
            )
        )
    return segments


def _normalize_chunk_decision(
    *,
    result: dict[str, object],
    topics: list[TopicNode],
    default_topic_id: str | None,
) -> dict[str, object]:
    valid_topic_ids = {topic.topic_id for topic in topics}
    allowed_decisions = {"link", "propose", "unclassified"}
    allowed_edge_types = {item.value for item in EdgeType}

    raw_decision = result.get("decision")
    decision = raw_decision if isinstance(raw_decision, str) and raw_decision in allowed_decisions else "unclassified"
    raw_topic_id = result.get("topic_id")
    topic_id = raw_topic_id if isinstance(raw_topic_id, str) and raw_topic_id in valid_topic_ids else None
    raw_reason = result.get("reason")
    reason = raw_reason.strip()[:80] if isinstance(raw_reason, str) and raw_reason.strip() else "未找到合适知识点"
    raw_guiding_question = result.get("guiding_question")
    guiding_question = raw_guiding_question.strip()[:60] if isinstance(raw_guiding_question, str) else ""
    raw_teaching_hint = result.get("teaching_hint")
    teaching_hint = raw_teaching_hint.strip()[:80] if isinstance(raw_teaching_hint, str) else ""

    raw_confidence = result.get("confidence", 0.0)
    try:
        confidence = max(0.0, min(1.0, float(raw_confidence)))
    except (TypeError, ValueError):
        confidence = 0.0

    proposed_topic: dict[str, object] | None = None
    if decision == "propose":
        raw_proposed_topic = result.get("proposed_topic")
        if isinstance(raw_proposed_topic, dict):
            raw_title = raw_proposed_topic.get("title")
            title = raw_title.strip()[:120] if isinstance(raw_title, str) else ""
            raw_summary = raw_proposed_topic.get("summary")
            summary = raw_summary.strip()[:500] if isinstance(raw_summary, str) else ""
            raw_parent_ids = raw_proposed_topic.get("parent_node_ids")
            parent_node_ids = [
                item for item in raw_parent_ids if isinstance(item, str) and item in valid_topic_ids
            ] if isinstance(raw_parent_ids, list) else []
            if not parent_node_ids and default_topic_id in valid_topic_ids:
                parent_node_ids = [default_topic_id]
            raw_edge_type = raw_proposed_topic.get("edge_type")
            edge_type = raw_edge_type if isinstance(raw_edge_type, str) and raw_edge_type in allowed_edge_types else "requires"
            if title:
                proposed_topic = {
                    "title": title,
                    "summary": summary,
                    "parent_node_ids": parent_node_ids,
                    "edge_type": edge_type,
                }

    if confidence < CLASSIFIED_THRESHOLD:
        decision = "unclassified"

    if decision == "link":
        if topic_id is None:
            decision = "unclassified"
    elif decision == "propose":
        topic_id = None
        if proposed_topic is None:
            decision = "unclassified"
    else:
        topic_id = None
        proposed_topic = None

    if decision == "unclassified":
        topic_id = None
        proposed_topic = None

    return {
        "decision": decision,
        "topic_id": topic_id,
        "confidence": confidence,
        "reason": reason,
        "proposed_topic": proposed_topic,
        "guiding_question": guiding_question,
        "teaching_hint": teaching_hint,
    }


def _tag_value(tags: list[str], prefix: str, default: str | None = None) -> str | None:
    for tag in tags:
        if tag.startswith(prefix):
            return tag.split(":", 1)[1]
    return default


def _node_subject_tag(tags: list[str]) -> str:
    raw = _tag_value(tags, "subject:")
    if raw is None:
        return "general"
    if raw in {"english", "chinese", "语文", "中文"}:
        return "language"
    return raw


def _detect_lang_from_chunks(chunks: list[TextUnit], topics: list[TopicNode]) -> str | None:
    sample = " ".join(c.text for c in chunks[:3]).lower()
    if "english" in sample or "英语" in sample or re.search(r"\b(present continuous|simple past)\b", sample):
        return "english"
    if "语文" in sample or "中文" in sample or "chinese" in sample or "阅读理解" in sample:
        return "chinese"
    return None


def _batch_organize_topic_tree(
    *,
    backend: "SessionBackend",
    candidates: list[str],
    subject: str,
    language_id: str | None = None,
) -> list[dict]:
    candidate_list = json.dumps(candidates, ensure_ascii=False)
    lang_hint = ""
    if subject == "language" and language_id:
        lang_hint = (
            f", 语言: {language_id}\n"
            "对于英语，按语法/词汇/发音/阅读/写作/功能句型分组。\n"
            "对于语文，按识字/古诗词/阅读理解/写作/语言知识点分组。\n"
        )
    system = (
        "你是知识结构组织助手。把一组概念关键词组织成树形层级。\n"
        f"学科: {subject}{lang_hint}"
        "规则:\n"
        "1. 中继节点标题必须以 [中继] 结尾\n"
        "2. 每个输入的关键词必须在树中出现一次\n"
        "3. 不要凭空创建不在列表中的叶子节点\n"
        "4. 同级概念平铺在同一中继下\n"
        "返回 JSON 树形数组。"
    )
    user = (
        f"概念关键词列表:\n{candidate_list}\n\n"
        "请组织为树形结构。中继节点标题加 [中继] 后缀。"
    )
    try:
        response = _llm_call(backend, 
            model=backend.llm_skill.model_name,
            messages=[{"role": "system", "content": system}, {"role": "user", "content": user}],
            temperature=0.3, timeout=600.0,
        )
        raw = (response.choices[0].message.content or "").strip()
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```$", "", raw)
        data = json.loads(raw)
        return data if isinstance(data, list) else []
    except Exception:
        return []


def _flatten_tree_prerequisites(tree: list[dict], topic_map: dict[str, str]) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}
    def _walk(nodes: list[dict]) -> list[str]:
        prev: str | None = None
        for node in nodes:
            raw_title = (node.get("title") or "").strip()
            clean = raw_title.replace("[中继]", "").strip()
            children = node.get("children") or []
            child_ids = _walk(children) if children else (
                [topic_map.get(clean, clean)] if clean in topic_map else []
            )
            for cid in child_ids:
                if prev and cid and cid != prev:
                    result.setdefault(cid, []).append(prev)
                prev = cid
        return child_ids if children else [topic_map.get(clean, clean)]
    _walk(tree)
    return result


def _create_relay_and_attach(
    *,
    backend: "SessionBackend",
    tree: list[dict],
    root_topic_id: str | None,
    topic_map: dict[str, str],
    subject: str,
    language_id: str | None,
    topics: list[TopicNode],
    parent_id: str | None = None,
) -> None:
    from src.services.resource_graph_curation import SUBJECT_PROFILES, _proposal_tags

    profile = SUBJECT_PROFILES.get(subject, SUBJECT_PROFILES["general"])
    for node in tree:
        raw_title = (node.get("title") or "").strip()
        if not raw_title:
            continue
        children = node.get("children") or []
        is_relay = "[中继]" in raw_title
        clean = raw_title.replace("[中继]", "").strip()

        if not children:
            # Leaf node: don't create here, GraphLocator will handle it
            continue

        # Relay node: create + approve
        pid = parent_id or root_topic_id
        if clean in topic_map:
            tid = topic_map[clean]
        else:
            try:
                proposal = backend.create_graph_proposal_from_resource(
                    title=clean,
                    summary=f"中继节点: {subject}学科",
                    tags=_proposal_tags(
                        subject=subject, facet=profile.default_facet,
                        language_id=language_id,
                    ),
                    parent_node_ids=[pid] if pid else [],
                    prerequisite_node_ids=[],
                    edge_type="part_of",
                    reason="batch tree relay node",
                )
                _st, _pr, tpc, _, _ = backend.approve_graph_proposal(
                    proposal_id=proposal.proposal_id, difficulty=1,
                )
                tid = tpc.topic_id
                topic_map[clean] = tid
                topics.append(tpc)
            except Exception:
                continue

        if children and tid:
            _create_relay_and_attach(
                backend=backend, tree=children,
                root_topic_id=root_topic_id, topic_map=topic_map,
                subject=subject, language_id=language_id,
                topics=topics, parent_id=tid,
            )


def _run_structured_ingestion(
    *,
    backend: "SessionBackend",
    record: ResourceRecord,
    chunks: list[TextUnit],
    topics: list[TopicNode],
    subject: str,
    language_id: str | None = None,
    progress_callback: Callable[[str, dict[str, object]], None] | None = None,
) -> list[ResourceSegment]:
    from src.services.resource_graph_curation import SUBJECT_PROFILES, _proposal_tags

    profile = SUBJECT_PROFILES.get(subject, SUBJECT_PROFILES["general"])

    from src.api.app import _is_cancelled as _ingestion_cancelled
    if _ingestion_cancelled(record.resource_id):
        return [_status_segment(resource_id=record.resource_id, sequence_index=0, status="parse_failed", reason="cancelled", locator={"kind": "structured"})]

    scan = _fast_document_scan(backend, chunks, subject, language_id)
    blocks = scan.get("blocks") or []
    full_text = "\n\n".join(c.text for c in chunks if c.text)
    _topic_block_types = {
        "topic_area", "grammar_point", "vocabulary_theme", "pronunciation",
        "reading", "functional_expression",
        "character_learning", "poetry", "reading_comprehension", "writing", "language_point",
    }
    topic_blocks = [b for b in blocks if b.get("type") in _topic_block_types]
    exercise_blocks = [b for b in blocks if b.get("type") == "exercise_only"]

    logger.info(
        "structured scan summary",
        extra={
            "resource_id": record.resource_id,
            "block_count": len(blocks),
            "topic_block_count": len(topic_blocks),
            "exercise_block_count": len(exercise_blocks),
            "block_types": [str((b or {}).get("type") or "") for b in blocks[:50]],
        },
    )

    if not topic_blocks and blocks:
        fallback_blocks = [
            b for b in blocks
            if str((b or {}).get("type") or "") not in {"word_list", "appendix", "exercise_only"}
        ]
        if fallback_blocks:
            topic_blocks = fallback_blocks
            for block in topic_blocks:
                block.setdefault("type", "topic_area")
            logger.warning(
                "no typed topic blocks from scan, using fallback blocks",
                extra={"resource_id": record.resource_id, "fallback_block_count": len(topic_blocks)},
            )

    if not topic_blocks and full_text.strip():
        topic_blocks = [{
            "label": "document",
            "summary": "full document fallback",
            "type": "topic_area",
            "start_marker": "",
            "end_marker": "",
            "source_text": full_text[:24000],
            "_whole_document_fallback": True,
        }]
        logger.warning(
            "scan returned no usable topic blocks, falling back to whole document extraction",
            extra={"resource_id": record.resource_id, "text_length": len(full_text)},
        )

    if not topic_blocks and not exercise_blocks:
        return [
            _status_segment(
                resource_id=record.resource_id, sequence_index=0,
                status="unclassified", reason="no topic or exercise blocks found",
                locator={"kind": "structured"},
            )
        ]

    topic_map = {t.title.strip(): t.topic_id for t in topics}
    seen_titles: set[str] = set()
    candidates: list[dict[str, object]] = []
    block_workers = min(6, max(1, len(topic_blocks)))

    if topic_blocks and backend.llm_skill.client.api_key:
        with ThreadPoolExecutor(max_workers=block_workers) as executor:
            futures = [
                executor.submit(
                    _process_topic_block,
                    backend,
                    block,
                    chunks,
                    subject,
                    language_id,
                    topic_map,
                )
                for block in topic_blocks
            ]
            for future in as_completed(futures):
                if _ingestion_cancelled(record.resource_id):
                    logger.info("structured ingestion cancelled during extraction", extra={"resource_id": record.resource_id})
                    break
                try:
                    block_result = future.result()
                except Exception:
                    logger.exception("topic block extraction failed", extra={"resource_id": record.resource_id})
                    continue

                block = block_result["block"]
                block_idx = -1
                for _bi, _orig in enumerate(topic_blocks):
                    if block is _orig:
                        block_idx = _bi
                        break
                if block_idx < 0:
                    continue
                block["topic_candidates"] = [item["title"] for item in block_result["topics"]]
                block["source_text"] = block_result["text"][:500]

                for candidate in block_result["topics"]:
                    title = candidate["title"]
                    desc = candidate["desc"]
                    title_key = _candidate_key(title)
                    if title_key in seen_titles:
                        _console_log("INGEST_QUEUE_DEDUP", {"skipped_title": title, "current_queue": sorted(seen_titles)})
                        continue
                    seen_titles.add(title_key)
                    _console_log(
                        "INGEST_QUEUE_STATE",
                        {
                            "enqueued_title": title,
                            "current_queue": sorted(seen_titles),
                        },
                    )
                    candidates.append({
                        "title": title,
                        "desc": desc,
                        "whole_document_fallback": bool(block.get("_whole_document_fallback")),
                        "block_idx": block_idx,
                    })

    _console_log("INGEST_CANDIDATE_BATCH", candidates)

    candidate_graph = _build_candidate_graph_from_extraction(
        record=record,
        candidates=candidates,
        topic_blocks=topic_blocks,
        subject=subject,
        language_id=language_id,
        chunks=chunks,
        backend=backend,
        scan=scan,
    )

    try:
        graph_json = candidate_graph.model_dump_json(indent=2, ensure_ascii=False)
    except Exception:
        graph_json = json.dumps(candidate_graph.model_dump(), ensure_ascii=False, default=str)
    backend.store_resource_candidate_graph(record.resource_id, graph_json)

    gross_node_count = len(candidate_graph.candidate_nodes)
    backend.update_resource_ingestion(
        record.resource_id,
        status="processing",
        error=None,
        extracted_node_count=gross_node_count,
    )
    _safe_progress(
        progress_callback,
        "chunk_classification_progress",
        {
            "resource_id": record.resource_id,
            "stage": "candidate_graph_built",
            "extracted_node_count": gross_node_count,
        },
    )
    logger.info(
        "candidate graph built",
        extra={
            "resource_id": record.resource_id,
            "gross_node_count": gross_node_count,
            "relay_count": len(candidate_graph.candidate_relays),
            "edge_count": len(candidate_graph.candidate_edges),
        },
    )

    # ── Review + Compile pipeline ──
    _set_pipeline_stage(progress_callback, record.resource_id, "review")
    try:
        from src.services.candidate_review import compile_review_result, review_candidate_graph

        review_result = review_candidate_graph(
            candidate_graph,
            existing_topics=topics,
            backend=backend,
        )
        _console_log("INGEST_REVIEW_RESULT", {
            "naming_passed": review_result.naming.passed,
            "logic_passed": review_result.logic.passed,
            "compilable_nodes": len(review_result.compilable_nodes),
            "dropped": review_result.dropped_temp_ids,
            "merge_map": review_result.merge_map,
        })

        _set_pipeline_stage(progress_callback, record.resource_id, "compiling")
        compile_result = compile_review_result(
            backend,
            review_result,
            candidate_graph,
            resource_id_override=record.resource_id,
        )
        _console_log("INGEST_COMPILE_RESULT", {
            "created_topics": compile_result.created_topic_count,
            "relays": compile_result.relay_count,
            "knowledge": compile_result.knowledge_count,
            "relinked_segments": compile_result.relinked_segment_count,
            "errors": compile_result.errors,
        })

        # Write actual created knowledge-node count (exclude relays)
        actual_node_count = compile_result.knowledge_count
        if compile_result.errors:
            logger.warning(
                "compile completed with errors",
                extra={
                    "resource_id": record.resource_id,
                    "created_topics": actual_node_count,
                    "planned_nodes": len(review_result.compilable_nodes),
                    "errors": compile_result.errors,
                },
            )
        backend.update_resource_ingestion(
            record.resource_id,
            status="processing",
            error=None,
            extracted_node_count=actual_node_count,
        )
        _safe_progress(
            progress_callback,
            "chunk_classification_progress",
            {
                "resource_id": record.resource_id,
                "stage": "compiled",
                "extracted_node_count": actual_node_count,
            },
        )

        # Store review result alongside candidate graph
        try:
            review_json = review_result.model_dump_json(indent=2, ensure_ascii=False)
            backend.store_resource_review_graph(record.resource_id, review_json)
            # Log summary to console for debugging; candidate_graph_json keeps the extraction output
            _console_log("INGEST_REVIEW_FULL", {
                "review_id": review_result.review_id,
                "naming_passed": review_result.naming.passed,
                "logic_passed": review_result.logic.passed,
                "compilable_count": len(review_result.compilable_nodes),
                "merge_map": review_result.merge_map,
                "dropped": review_result.dropped_temp_ids,
            })
        except Exception:
            pass
    except Exception as exc:
        error_msg = str(exc)[:200]
        logger.exception("review/compile pipeline failed", extra={"resource_id": record.resource_id, "error": error_msg})
        _safe_progress(
            progress_callback,
            "chunk_classification_progress",
            {
                "resource_id": record.resource_id,
                "stage": "review_failed",
                "error": error_msg,
            },
        )
        # Propagate failure: mark ingestion as failed so outer layer and frontend see the error
        try:
            backend.update_resource_ingestion(record.resource_id, status="failed", error=f"review_compile: {error_msg}")
        except Exception:
            pass
        _set_pipeline_stage(progress_callback, record.resource_id, "failed")

    if exercise_blocks:
        _process_exercise_blocks(
            backend=backend, blocks=exercise_blocks, chunks=chunks,
            topics=topics, topic_map=topic_map,
            record=record, subject=subject, language_id=language_id,
            profile=profile, _proposal_tags=_proposal_tags,
        )

    return _build_segments_from_scan(
        scan=scan, record=record, topics=topics, topic_map=topic_map, chunks=chunks,
    )


def _set_pipeline_stage(
    progress_callback: Callable[[str, dict[str, object]], None] | None,
    resource_id: str,
    stage: str,
) -> None:
    try:
        from src.api.app import _set_pipeline_stage as _api_set_pipeline_stage
        _api_set_pipeline_stage(resource_id, stage)
    except Exception:
        pass


# ── Phase 1 helpers ──

def _get_model_context_limit(backend: "SessionBackend") -> int:
    import os as _os, json as _json
    from pathlib import Path as _Path
    data_root = _os.getenv("DATA_ROOT", "./data")
    settings_path = _Path(data_root) / "llm_settings.json"
    settings: dict = {}
    if settings_path.exists():
        try:
            settings = _json.loads(settings_path.read_text(encoding="utf-8"))
        except Exception:
            pass
    mode = str(settings.get("mode") or "remote")
    model_catalog = settings.get("llm_models") if isinstance(settings.get("llm_models"), dict) else {}
    if mode == "local":
        model_name = str((settings.get("local") or {}).get("model") or (settings.get("local") or {}).get("model_path") or "")
    else:
        model_name = str((settings.get("remote") or {}).get("model") or "")
    if model_catalog:
        model_key = str(model_name).strip()
        for key, model_def in model_catalog.items():
            if not isinstance(model_def, dict):
                continue
            candidates = {
                str(key).strip(),
                str(model_def.get("label") or "").strip(),
                str(model_def.get("repo_id") or "").strip(),
            }
            if model_key not in candidates:
                continue
            raw_limit = model_def.get("context_limit")
            try:
                parsed_limit = int(str(raw_limit).strip())
                if parsed_limit > 0:
                    return parsed_limit
            except (TypeError, ValueError):
                pass
    model_lower = model_name.lower()
    if "deepseek" in model_lower:
        return 1000000
    if "qwen" in model_lower:
        return 131072
    if "gemma" in model_lower:
        return 32768
    return 32768


def _fast_document_scan(
    backend: "SessionBackend",
    chunks: list[TextUnit],
    subject: str,
    language_id: str | None = None,
) -> dict:
    full_text = "\n\n".join(c.text for c in chunks if c.text)
    context_limit = _get_model_context_limit(backend)
    max_tokens = int(context_limit * 0.8)
    text_tokens = _estimate_tokens(full_text)
    _console_log(
        "INGEST_SCAN_TOKEN_SUMMARY",
        {
            "total_text_chars": len(full_text),
            "total_estimated_tokens": text_tokens,
            "context_limit": context_limit,
            "scan_max_tokens": max_tokens,
            "chunk_count": len(chunks),
        },
    )
    if text_tokens <= max_tokens:
        return _fast_document_scan_single(backend, full_text, subject, language_id)

    # Split text into chunks fitting within token limit
    window_chars = int(len(full_text) * max_tokens / max(text_tokens, 1))
    overlap = window_chars // 6
    blocks_all = []
    start = 0
    while start < len(full_text):
        end = min(start + window_chars, len(full_text))
        part = full_text[start:end]
        _console_log(
            "INGEST_SCAN_WINDOW",
            {
                "window_start": start,
                "window_end": end,
                "window_chars": len(part),
                "window_estimated_tokens": _estimate_tokens(part),
            },
        )
        partial_scan = _fast_document_scan_single(backend, part, subject, language_id)
        blocks = partial_scan.get("blocks") or []
        for b in blocks:
            b["_offset"] = start
        blocks_all.extend(blocks)
        if end >= len(full_text):
            break
        start = end - overlap
    return {"doc_type": "textbook", "blocks": blocks_all}


def _attach_block_source_texts(scan_text: str, blocks: list[dict]) -> list[dict]:
    if not scan_text.strip():
        return blocks

    normalized_text = scan_text
    cursor = 0
    block_count = len(blocks)
    for index, block in enumerate(blocks):
        if not isinstance(block, dict):
            continue
        start_marker = str(block.get("start_marker") or "").strip()
        end_marker = str(block.get("end_marker") or "").strip()
        label = str(block.get("label") or "").strip()
        summary = str(block.get("summary") or "").strip()

        start_pos = cursor
        for marker in (start_marker, label, summary):
            if not marker:
                continue
            found = normalized_text.find(marker, cursor)
            if found != -1:
                start_pos = found
                break

        end_pos = len(normalized_text)
        if end_marker:
            found_end = normalized_text.find(end_marker, start_pos + max(len(start_marker), 1))
            if found_end != -1:
                end_pos = found_end
        if end_pos == len(normalized_text) and index + 1 < block_count:
            next_block = blocks[index + 1]
            if isinstance(next_block, dict):
                next_markers = [
                    str(next_block.get("start_marker") or "").strip(),
                    str(next_block.get("label") or "").strip(),
                    str(next_block.get("summary") or "").strip(),
                ]
                for marker in next_markers:
                    if not marker:
                        continue
                    found_next = normalized_text.find(marker, start_pos + max(len(start_marker), 1))
                    if found_next != -1:
                        end_pos = found_next
                        break

        snippet = normalized_text[start_pos:end_pos].strip()
        if not snippet:
            snippet = summary or label or normalized_text[max(0, start_pos): min(len(normalized_text), start_pos + 2000)].strip()
        block["source_text"] = snippet[:12000]
        cursor = max(cursor, end_pos if end_pos > start_pos else start_pos)
    return blocks


def _fast_document_scan_single(
    backend: "SessionBackend",
    full_text: str,
    subject: str,
    language_id: str | None = None,
) -> dict:

    if subject == "language" and language_id == "english":
        system = (
            "你是英语教材分析助手。通读以下英语教学文档全文。\n"
            "任务:\n"
            "1. 判断文档类型: textbook | workbook | examination | vocabulary_list | phonics_chart\n"
            "2. 识别单元结构 (Unit 1, Unit 2...)\n"
            "3. 对每个单元提取以下类型的知识块:\n"
            '   - grammar_point: 语法知识点 (如"现在进行时"、"情态动词 must")\n'
            '     → type="grammar_point", topic_candidates填语法点名称\n'
            '   - vocabulary_theme: 词汇主题 (如"School Rules"、"Weather")\n'
            '     → type="vocabulary_theme", topic_candidates填主题分类名\n'
            '   - pronunciation: 发音规则 (如"a_e /eɪ/"、"aw /ɔː/")\n'
            '     → type="pronunciation", topic_candidates填发音规则\n'
            '   - reading: 阅读课文 (如有明确课文标题)\n'
            '     → type="reading", topic_candidates填课文标题或主题\n'
            '   - functional_expression: 功能句型 (如"Making Requests")\n'
            '     → type="functional_expression", topic_candidates填功能描述\n'
            "4. 单词表/词汇表标记为 type=word_list, 不提取 topic_candidates\n"
            "5. 练习题标记为 type=exercise_only\n"
            "返回 JSON: {\"doc_type\":\"...\",\"blocks\":[{\"label\":\"...\",\"summary\":\"...\","
            '"type":"grammar_point|vocabulary_theme|pronunciation|reading|functional_expression|'
            'word_list|exercise_only|fuzzy|appendix","start_marker":"...","end_marker":"..."}]'
        )
    elif subject == "language" and language_id == "chinese":
        system = (
            "你是语文教材分析助手。通读以下语文教学文档全文。\n"
            "任务:\n"
            "1. 判断文档类型: textbook | workbook | examination | poetry_collection | character_practice\n"
            "2. 识别课文结构 (第X课、第X单元...)\n"
            "3. 对每课/每单元提取:\n"
            '   - character_learning: 生字/识字点\n'
            '     → type="character_learning", topic_candidates填识字主题 (如"会认字:山水火")\n'
            '   - poetry: 古诗词\n'
            '     → type="poetry", topic_candidates填诗词标题\n'
            '   - reading_comprehension: 课文阅读理解\n'
            '     → type="reading_comprehension", topic_candidates填课文标题或理解主题\n'
            '   - writing: 写作练习\n'
            '     → type="writing", topic_candidates填写作类型 (如"看图写话"、"日记")\n'
            '   - language_point: 语言知识点\n'
            '     → type="language_point", topic_candidates填知识点 (如"比喻"、"排比")\n'
            "4. 生字表/写字表标记为 type=word_list\n"
            "5. 练习题标记为 type=exercise_only\n"
            "返回 JSON: {\"doc_type\":\"...\",\"blocks\":[{\"label\":\"...\",\"summary\":\"...\","
            '"type":"character_learning|poetry|reading_comprehension|writing|language_point|'
            'word_list|exercise_only|fuzzy|appendix","start_marker":"...","end_marker":"..."}]'
        )
    else:
        system = (
            "你是教学文档分析助手。通读以下文档全文，返回结构化分析。\n"
            f"学科: {subject}\n"
            "输出 JSON：{\"doc_type\":\"...\",\"blocks\":[{\"label\":\"...\",\"summary\":\"...\","
            '"type":"topic_area|exercise_only|fuzzy|word_list|appendix","start_marker":"...","end_marker":"..."}]'
        )
    user = f"文档全文:\n```text\n{full_text}\n```\n请分析并返回 JSON。"
    _console_log(
        "INGEST_SCAN_PROMPT",
        {
            "subject": subject,
            "language_id": language_id,
            "full_text_chars": len(full_text),
            "full_text_estimated_tokens": _estimate_tokens(full_text),
            "system_prompt": system,
            "user_prompt": "文档全文: <omitted resource text>\\n请分析并返回 JSON。",
        },
    )

    fallback = {"doc_type": "unknown", "blocks": []}
    try:
        response = _llm_call(backend, 
            model=backend.llm_skill.model_name,
            messages=[{"role": "system", "content": system}, {"role": "user", "content": user}],
            temperature=0.3, timeout=600.0,
        )
        raw = (response.choices[0].message.content or "").strip()
        _console_log("INGEST_SCAN_RESPONSE", raw)
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```$", "", raw)
        data = json.loads(raw)
        if isinstance(data, dict):
            blocks = data.get("blocks") or data.get("sections") or []
            if isinstance(blocks, list):
                data["blocks"] = _attach_block_source_texts(full_text, [block for block in blocks if isinstance(block, dict)])
            return data
        return fallback
    except Exception:
        return fallback


# ── Phase 2 helpers ──

def _process_topic_block(
    backend: "SessionBackend",
    block: dict,
    chunks: list[TextUnit],
    subject: str,
    language_id: str | None,
    topic_map: dict[str, str],
) -> dict[str, object]:
    block_text = str(block.get("source_text") or "").strip() or _extract_block_text(block, chunks)
    _console_log(
        "INGEST_TOPIC_BLOCK_INPUT",
        {
            "label": block.get("label"),
            "type": block.get("type"),
            "summary": block.get("summary"),
            "block_text_chars": len(block_text),
            "block_text_estimated_tokens": _estimate_tokens(block_text),
        },
    )
    topics_data = _extract_topics_from_block_text(backend, block_text, subject, language_id)
    results: list[dict[str, str]] = []
    for td in topics_data:
        title = (td.get("title") or "").strip()
        desc = (td.get("desc") or "").strip()
        if title:
            results.append({"title": title, "desc": desc})
    if not results:
        fallback_title = str(block.get("label") or "").strip()
        fallback_desc = str(block.get("summary") or "").strip()
        if fallback_title and fallback_title.lower() not in {"document", "words", "word list", "appendix"}:
            results.append({"title": fallback_title[:120], "desc": fallback_desc[:200]})
    _console_log("INGEST_TOPIC_BLOCK_OUTPUT", {"label": block.get("label"), "topics": results})
    return {"block": block, "text": block_text, "topics": results}


def _candidate_key(title: str) -> str:
    normalized = re.sub(r"\s+", " ", title.strip().lower())
    return normalized[:200]


def _resolve_parent_ids(parent_ids: list[str], topic_map: dict[str, str]) -> list[str]:
    resolved: list[str] = []
    for parent_id in parent_ids:
        if parent_id in topic_map.values():
            resolved.append(parent_id)
            continue
        if parent_id.startswith("relay_"):
            relay_title = parent_id[len("relay_"):].strip()
            mapped = topic_map.get(relay_title)
            if mapped:
                resolved.append(mapped)
                continue
        mapped = topic_map.get(parent_id)
        if mapped:
            resolved.append(mapped)
    return list(dict.fromkeys(resolved))


def _insert_topic_candidate(
    *,
    backend: "SessionBackend",
    locator,
    topic_map: dict[str, str],
    topics: list[TopicNode],
    title: str,
    desc: str,
    root_topic_id: str | None,
    subject: str,
    language_id: str | None,
    profile: object,
    proposal_tags_builder,
    placement: dict[str, object] | None = None,
    use_root_shortcut: bool = False,
) -> bool:
    existing = _link_to_topic(title, topics, topic_map)
    if existing:
        _console_log("INGEST_INSERT_SKIP_EXISTING", {"title": title, "matched_topic_id": existing})
        return False

    pos_exists = False
    pos_node_id = None
    pos_parent_ids: list[str] = []
    pos_successor_ids: list[str] = []
    pos_reason = ""
    if placement is not None:
        pos_exists = bool(placement.get("exists", False))
        pos_node_id = placement.get("node_id")
        pos_parent_ids = [pid for pid in (placement.get("parent_ids") or []) if isinstance(pid, str)]
        pos_successor_ids = [pid for pid in (placement.get("successor_ids") or []) if isinstance(pid, str)]
        pos_reason = str(placement.get("reason", ""))
    else:
        pos = locator.locate(title, desc)
        pos_exists = pos.exists
        pos_node_id = pos.node_id
        pos_parent_ids = list(pos.parent_ids)
        pos_successor_ids = list(pos.successor_ids)
        pos_reason = pos.reason

    if pos_exists and pos_node_id:
        topic_map[title] = pos_node_id
        _console_log("INGEST_INSERT_REUSED", {"title": title, "node_id": pos_node_id, "reason": pos_reason})
        return False

    parents = _resolve_parent_ids(pos_parent_ids, topic_map)
    if use_root_shortcut and root_topic_id:
        parents = [root_topic_id]
    elif not parents and root_topic_id:
        parents = [root_topic_id]
    _console_log(
        "INGEST_INSERT_ACTION",
        {
            "title": title,
            "desc": desc[:200],
            "parent_ids": parents,
            "locator_reason": pos_reason,
        },
    )
    try:
        proposal = backend.create_graph_proposal_from_resource(
            title=title,
            summary=desc[:200],
            tags=proposal_tags_builder(
                subject=subject,
                facet=profile.default_facet,
                language_id=language_id,
            ),
            parent_node_ids=parents,
            prerequisite_node_ids=_resolve_parent_ids(pos_successor_ids, topic_map),
            edge_type="part_of" if parents else "requires",
            reason=pos_reason[:80],
        )
        _st, _pr, topic, _, _ = backend.approve_graph_proposal(
            proposal_id=proposal.proposal_id,
            difficulty=1,
        )
        _console_log(
            "INGEST_INSERT_RESULT",
            {
                "title": title,
                "proposal_id": proposal.proposal_id,
                "topic_id": topic.topic_id,
            },
        )
        topic_map[title] = topic.topic_id
        topics.append(topic)
        locator.refresh()
        return True
    except Exception:
        logger.exception(
            "topic insertion failed",
            extra={"title": title},
        )
        return False


def _extract_block_text(block: dict, chunks: list[TextUnit]) -> str:
    start = (block.get("start_marker") or "").lower()
    end = (block.get("end_marker") or "").lower()
    if not start:
        return " ".join(c.text for c in chunks[:3])

    capturing = False
    parts: list[str] = []
    for chunk in chunks:
        text = chunk.text or ""
        lower = text.lower()
        if start in lower:
            capturing = True
        if capturing:
            parts.append(text)
        if end and end in lower:
            break
    return "\n".join(parts) if parts else " ".join(c.text for c in chunks[:3])


def _extract_topics_from_block_text(
    backend: "SessionBackend",
    block_text: str,
    subject: str,
    language_id: str | None = None,
) -> list[dict]:
    lang_hint = f", 语言: {language_id}" if language_id else ""
    system = (
        "你是教学知识点提取助手。从段落提取可作知识图谱节点的知识点。\n"
        "只提取可教学的知识点（定理、定律、公式、概念、语法点）。\n"
        "不要提取实验过程、课堂活动、练习指令、方法论标题。\n"
        f"学科: {subject}{lang_hint}\n"
        '返回 JSON 数组: [{"title":"知识点名","desc":"一句话说明"}]'
    )
    user = f"段落:\n```text\n{block_text[:3000]}\n```\n请提取知识点。"
    _console_log(
        "INGEST_TOPIC_EXTRACT_PROMPT",
        {
            "subject": subject,
            "language_id": language_id,
            "block_text_chars": len(block_text),
            "block_text_estimated_tokens": _estimate_tokens(block_text),
            "system_prompt": system,
            "user_prompt": "段落: <omitted resource text>\\n请提取知识点。",
        },
    )
    try:
        response = _llm_call(backend, 
            model=backend.llm_skill.model_name,
            messages=[{"role": "system", "content": system}, {"role": "user", "content": user}],
            temperature=0.3, timeout=600.0,
        )
        raw = (response.choices[0].message.content or "").strip()
        _console_log("INGEST_TOPIC_EXTRACT_RESPONSE", raw)
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```$", "", raw)
        data = json.loads(raw)
        return [item for item in data if isinstance(item, dict)] if isinstance(data, list) else []
    except Exception:
        return []


def _find_subject_root_id(topics: list[TopicNode], subject: str, language_id: str | None) -> str | None:
    # Find root by subject tag + facet:root, not by generic title
    for t in topics:
        tags = set(t.tags)
        if "facet:root" not in tags:
            continue
        if f"subject:{subject}" not in tags:
            continue
        if subject == "language" and language_id and f"language:{language_id}" not in tags:
            continue
        return t.topic_id
    # Fallback: any facet:root with matching title
    for t in topics:
        if "facet:root" in set(t.tags) and t.title.strip().lower() == subject.lower():
            return t.topic_id
    return None


def _link_to_topic(title: str, topics: list[TopicNode], topic_map: dict[str, str]) -> str | None:
    clean = title.strip()
    if clean in topic_map:
        return topic_map[clean]
    for t in topics:
        if t.title.strip() == clean:
            topic_map[clean] = t.topic_id
            return t.topic_id
    return None


def _process_exercise_blocks(
    *,
    backend: "SessionBackend",
    blocks: list[dict],
    chunks: list[TextUnit],
    topics: list[TopicNode],
    topic_map: dict[str, str],
    record: ResourceRecord,
    subject: str,
    language_id: str | None,
    profile: object,
    _proposal_tags: object,
) -> None:
    pass  # Exercises processed in segment builder


def _build_segments_from_scan(
    *,
    scan: dict,
    record: ResourceRecord,
    topics: list[TopicNode],
    topic_map: dict[str, str],
    chunks: list[TextUnit],
) -> list[ResourceSegment]:
    blocks = scan.get("blocks") or []
    segments: list[ResourceSegment] = []
    seg_index = 0

    for block in blocks:
        block_text = str(block.get("source_text") or "").strip() or _extract_block_text(block, chunks)
        if not block_text:
            block_text = block.get("summary", "")
        btype = block.get("type", "unknown")
        label = block.get("label", btype)
        candidate_titles = [str(item).strip() for item in (block.get("topic_candidates") or []) if str(item).strip()]

        seg_index += 1
        if btype in {"word_list", "appendix"}:
            segments.append(
                ResourceSegment(
                    segment_id=_segment_id(record.resource_id, seg_index),
                    start_ms=0, end_ms=None, label=btype,
                    status="unclassified", sequence_index=seg_index,
                    text=block.get("summary", block_text)[:500],
                    locator={"kind": "structured", "section_type": btype},
                    topic_id=None, proposal_id=None,
                    decision="unclassified", confidence=0.0,
                    reason=f"{btype} section skipped",
                )
            )
        else:
            matched_tid = None
            preferred_titles = candidate_titles or [label, block.get("summary", "")]
            for preferred_title in preferred_titles:
                matched_tid = _link_to_topic(preferred_title, topics, topic_map)
                if matched_tid:
                    break
            if matched_tid is None:
                for title, tid in topic_map.items():
                    if title.lower() in (label or "").lower() or title.lower() in (block.get("summary", "") or "").lower():
                        matched_tid = tid
                        break
            segments.append(
                ResourceSegment(
                    segment_id=_segment_id(record.resource_id, seg_index),
                    start_ms=0, end_ms=None, label="chunk",
                    status="classified" if matched_tid else "unclassified",
                    sequence_index=seg_index,
                    text=block.get("summary", label)[:500],
                    locator={"kind": "structured", "section_type": btype},
                    topic_id=matched_tid,
                    proposal_id=None,
                    decision="link" if matched_tid else "unclassified",
                    confidence=0.85,
                    reason="via structured scan" if matched_tid else "no topic extracted",
                )
            )
    return segments


def _build_candidate_graph_from_extraction(
    *,
    record: ResourceRecord,
    candidates: list[dict[str, object]],
    topic_blocks: list[dict],
    subject: str,
    language_id: str | None,
    chunks: list[TextUnit],
    backend: "SessionBackend",
    scan: dict,
) -> CandidateGraph:
    import datetime as _dt
    import time as _time

    now_iso = _dt.datetime.now(_dt.timezone.utc).isoformat()
    full_text = "\n\n".join(c.text for c in chunks if c.text)

    doc_meta = DocumentMeta(
        resource_id=record.resource_id,
        resource_name=record.resource_name,
        media_type=record.media_type,
        original_filename=record.original_filename,
        subject=subject,
        language_id=language_id,
        doc_type=str(scan.get("doc_type", "")),
        total_chunks=len(chunks),
        total_chars=len(full_text),
        extraction_timestamp=now_iso,
    )

    source_segments: list[SourceSegment] = []
    seg_index = 0
    for block in topic_blocks:
        seg_index += 1
        block_text = str(block.get("source_text") or block.get("summary", ""))
        source_segments.append(
            SourceSegment(
                segment_id=_segment_id(record.resource_id, seg_index),
                segment_index=seg_index,
                locator={
                    "kind": "structured",
                    "section_type": block.get("type", "unknown"),
                    "block_label": str(block.get("label", "")),
                },
                text=block_text[:2000],
                status="unclassified",
            )
        )

    candidate_nodes: list[CandidateNode] = []
    title_to_idx: dict[str, int] = {}
    for idx, candidate in enumerate(candidates, start=1):
        title = str(candidate["title"]).strip()
        desc = str(candidate["desc"]).strip()
        temp_id = f"cand_{idx:03d}"
        title_to_idx[title] = idx

        block_idx_val = int(candidate.get("block_idx", idx - 1)) if isinstance(candidate.get("block_idx"), (int, float)) else idx - 1
        block_idx = max(0, min(block_idx_val, len(topic_blocks) - 1))
        block = topic_blocks[block_idx] if topic_blocks else {}
        block_label = str(block.get("label") or "")

        refs: list[SourceRef] = []
        if block_idx < len(source_segments):
            refs.append(
                SourceRef(
                    segment_id=source_segments[block_idx].segment_id,
                    excerpt=title[:120],
                    relevance=0.9,
                )
            )

        candidate_nodes.append(
            CandidateNode(
                temp_id=temp_id,
                title=title,
                normalized_title=title,
                subject=subject,
                language_id=language_id,
                facet="general",
                summary=desc[:200],
                confidence=0.85,
                extraction_label=block_label or None,
                source_segment_refs=refs,
            )
        )

    relays: list[CandidateRelay] = []
    edges: list[CandidateEdge] = []
    block_to_nodes: dict[int, list[str]] = {}
    for idx, candidate in enumerate(candidates, start=1):
        b_idx = int(candidate.get("block_idx", idx - 1)) if isinstance(candidate.get("block_idx"), (int, float)) else idx - 1
        b_idx = max(0, min(b_idx, len(topic_blocks) - 1))
        block_to_nodes.setdefault(b_idx, []).append(f"cand_{idx:03d}")

    relay_idx = 0
    for b_idx, node_ids in block_to_nodes.items():
        if len(node_ids) < 2:
            continue
        relay_idx += 1
        relay_id = f"relay_{relay_idx:03d}"
        block = topic_blocks[b_idx] if b_idx < len(topic_blocks) else {}
        relay_title = str(block.get("label") or block.get("summary") or f"Section {relay_idx}")

        relays.append(
            CandidateRelay(
                relay_id=relay_id,
                title=relay_title,
                subject=subject,
                language_id=language_id,
                facet="topic_area",
                children_temp_ids=list(node_ids),
                confidence=0.8,
                grouping_rationale=f"Nodes extracted from same document section: {relay_title}",
            )
        )

        for node_id in node_ids:
            edges.append(
                CandidateEdge(
                    edge_id=f"edge_{len(edges)+1:03d}",
                    source_temp_id=node_id,
                    target_temp_id=relay_id,
                    edge_type="part_of",
                    edge_kind="parent_child",
                    confidence=0.85,
                    rationale=f"Belongs to {relay_title}",
                )
            )

    model_name = getattr(getattr(backend, "llm_skill", None), "model_name", "") or ""
    diag = ExtractionDiagnostics(
        extraction_model=model_name,
        pipeline="structured_scan",
        node_count=len(candidate_nodes),
        relay_count=len(relays),
        edge_count=len(edges),
        merge_group_count=0,
        segment_count=len(source_segments),
    )

    return CandidateGraph(
        extraction_id=f"extr_{int(_time.time()*1000)}_{os.urandom(3).hex()}",
        document=doc_meta,
        candidate_nodes=candidate_nodes,
        candidate_relays=relays,
        candidate_edges=edges,
        source_segments=source_segments,
        diagnostics=diag,
    )


def _status_segment(
    *,
    resource_id: str,
    sequence_index: int,
    status: str,
    reason: str,
    locator: dict[str, object],
) -> ResourceSegment:
    return ResourceSegment(
        segment_id=_segment_id(resource_id, sequence_index),
        start_ms=0,
        end_ms=None,
        label="document",
        status=status,
        sequence_index=sequence_index,
        text="",
        locator=locator,
        topic_id=None,
        decision="unclassified",
        confidence=0.0,
        reason=reason[:80],
        guiding_question=None,
        teaching_hint=None,
    )


def _segment_id(resource_id: str, sequence_index: int) -> str:
    return f"seg_{resource_id}_{sequence_index:04d}_{uuid.uuid4().hex[:6]}"
